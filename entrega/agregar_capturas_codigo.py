from pathlib import Path
from copy import deepcopy
from shutil import copy2
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
R=Path(r'C:\Users\Tecsup\AndroidStudioProjects\Lab09')
source=Path(r'C:\Users\Tecsup\Downloads\GLAB-S09-BPAREJA-2026-01.docx')
base=R/'tmp/word-antes-capturas.docx'
if not base.exists(): copy2(source,base)
d=Document(base)
original=Document(R/'tmp/GLAB-S09-original.docx')
# Mantener exactamente los elementos de la carátula original del profesor.
for i in range(20):
 old=d.paragraphs[i]._p
 old.getparent().replace(old,deepcopy(original.paragraphs[i]._p))

def add_after(anchor,text=''):
 el=OxmlElement('w:p'); anchor._p.addnext(el)
 p=Paragraph(el,anchor._parent)
 p.paragraph_format.space_after=Pt(5)
 if text:
  run=p.add_run(text);run.font.name='Arial';run.font.size=Pt(9)
 return p

def add_capture(anchor,name,caption):
 title=add_after(anchor,'Código implementado: '+caption)
 title.paragraph_format.keep_with_next=True
 p=add_after(title)
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.add_run().add_picture(str(R/'evidencias/codigo'/name),width=Inches(5.75))
 p.paragraph_format.keep_together=True
 return p

def find(starts):
 return next(p for p in d.paragraphs if p.text.startswith(starts))

add_capture(find('2.2. Agregue permisos'),'01-permiso-internet.png','permiso INTERNET en AndroidManifest.xml')
add_capture(find('Implementación realizada:'),'02-dependencias.png','dependencias de app/build.gradle.kts')
add_capture(find('El modelo se implementó'),'03-modelo-post.png','PostModel y nombres de los campos JSON')
add_capture(find('4.1. Cree una interfaz'),'04-servicio-post.png','PostApiService con GET y Path')
add_capture(find('Adaptación MVVM:'),'05-retrofit.png','clientes Retrofit e instancias de los servicios')
add_capture(find('5.3.1.'),'06-navegacion.png','ProgPrincipal9 con Scaffold y navegación inferior')
add_capture(find('5.3.2.'),'07-rutas.png','NavHost y rutas del listado y detalle')
add_capture(find('El archivo ui/moduloPosts.kt'),'08-publicaciones.png','PostsScreen y LazyColumn')
# Captura después de la respuesta MVVM, antes de continuar con ejecución.
anchor=Paragraph(OxmlElement('w:p'),d._body)
d.tables[20]._tbl.addnext(anchor._p)
add_capture(anchor,'09-mvvm.png','LabViewModel con estados de carga, resultado y error')
for table_index,name,caption in [(21,'10-recetas.png','búsqueda y tarjetas del catálogo de recetas'),(22,'11-anfibios.png','pantalla de anfibios e imágenes con Coil')]:
 cell=d.tables[table_index].cell(0,0)
 p=cell.add_paragraph('Código implementado: '+caption)
 p.paragraph_format.keep_with_next=True
 for run in p.runs: run.font.size=Pt(9);run.font.name='Arial'
 p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.add_run().add_picture(str(R/'evidencias/codigo'/name),width=Inches(5.6))

out=R/'entrega/GLAB-S09-BPAREJA-2026-01-completado.docx'
d.save(out)
# Comprobación estructural de la carátula antes de sustituir el archivo.
new=Document(out)
assert all(etree.tostring(new.paragraphs[i]._p)==etree.tostring(original.paragraphs[i]._p) for i in range(20))
copy2(out,source)
print('Capturas añadidas: 11. Carátula original conservada. Total de imágenes:',len(new.inline_shapes))
