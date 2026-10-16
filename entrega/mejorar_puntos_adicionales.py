from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

R=Path(__file__).resolve().parent.parent
source=Path(r'C:\Users\Tecsup\Downloads\GLAB-S09-BPAREJA-2026-01.docx')
backup=R/'tmp/word-antes-mejora-anfibios.docx'
if not backup.exists(): copy2(source,backup)
d=Document(backup)
cell=d.tables[22].cell(0,0)
cell.text=''
for e in d.tables[22]._tbl.xpath('./w:tr/w:trPr/w:trHeight | ./w:tr/w:trPr/w:cantSplit'):
    e.getparent().remove(e)

def para(text='',bold=False):
    p=cell.paragraphs[0] if len(cell.paragraphs)==1 and not cell.paragraphs[0].text and len(cell.paragraphs[0].runs)==0 else cell.add_paragraph()
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.line_spacing=1.08
    r=p.add_run(text); r.bold=bold;r.font.name='Arial';r.font.size=Pt(10)
    return p

def pictures(names,caption,width):
    p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    for i,name in enumerate(names):
        if i:p.add_run('   ')
        p.add_run().add_picture(str(R/name),width=Inches(width))
    p=para(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.italic=True;r.font.size=Pt(9)

para('Reto adicional implementado y probado',True)
para('Implementé un catálogo interactivo de anfibios que consulta seis especies mediante GET https://android-kotlin-fun-mars-server.appspot.com/amphibians. La aplicación permite explorar las especies, buscar coincidencias, filtrar por tipo y abrir una ficha con imagen y descripción en español.')
para('Funcionamiento',True)
for t in [
'1. Abrir la pestaña Anfibios. Se consultan los datos y se muestran las seis especies con su imagen, nombre, tipo y descripción.',
'2. Seleccionar Rana, Salamandra o Sapo para filtrar el catálogo. El contador indica cuántos resultados coinciden. Todos restaura los tipos disponibles.',
'3. Escribir en Buscar especie. La búsqueda combina el texto con el filtro seleccionado. En la prueba, Rana mostró dos resultados y la búsqueda azules redujo la lista a una especie.',
'4. Pulsar una tarjeta o Ver ficha de la especie para abrir el detalle. Volver al catálogo conserva la búsqueda y el filtro.',
'5. Pulsar Actualizar para consultar nuevamente la API. Borrar elimina el texto; Limpiar filtros restablece el catálogo cuando no hay coincidencias.'
]:para(t)
para('Implementación y validación',True)
para('Retrofit obtiene los datos desde el repositorio; el ViewModel ejecuta la consulta con corrutinas y expone los estados de carga, éxito y error. Compose aplica la búsqueda y el filtro sobre la respuesta recibida. rememberSaveable conserva la selección y la ficha abierta; Coil carga las imágenes. La compilación terminó correctamente en Android Studio.')
para('Se verificaron en el emulador: carga de seis especies, filtro de dos ranas, búsqueda con un resultado, apertura y retorno del detalle, actualización de datos, búsqueda sin coincidencias y limpieza del catálogo. El recorrido real se grabó en un video específico del reto.')
p=para('Video del reto adicional: ')
link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to('https://github.com/leonel10164/Lab09/blob/main/evidencias/anfibios/reto-anfibios.mp4',RT.HYPERLINK,is_external=True))
r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');rp.append(color);r.append(rp);t=OxmlElement('w:t');t.text='Abrir demostración del funcionamiento';r.append(t);link.append(r);p._p.append(link)
pictures(['evidencias/anfibios/01-catalogo.png','evidencias/anfibios/02-filtro-ranas.png'],'Evidencia del emulador: catálogo completo con seis especies y filtro Rana con dos resultados.',2.25)
pictures(['evidencias/anfibios/04-detalle.png'],'Evidencia del emulador: ficha de la rana de patas azules abierta desde la búsqueda.',2.25)
pictures(['evidencias/android-studio/15-anfibios-busqueda-filtros.png'],'Captura real de Android Studio: búsqueda, filtros y actualización del catálogo.',5.6)
pictures(['evidencias/android-studio/14-anfibios-detalle.png'],'Captura real de Android Studio: ficha de la especie y retorno al catálogo.',5.6)
out=R/'entrega/GLAB-S09-BPAREJA-2026-01-completado.docx'
d.save(out)
try:
    copy2(out,source)
    print('Documento original actualizado.')
except PermissionError:
    print('El original sigue abierto. La copia de entrega está actualizada.')
copy2(R/'app/build/outputs/apk/debug/app-debug.apk',R/'entrega/Lab09.apk')
print('Copia de entrega y APK actualizados.')
